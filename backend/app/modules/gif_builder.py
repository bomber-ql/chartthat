"""M4b 动效产物 GIF 模块（SPEC §5.6a / §6a，v1.11 引入 / v1.12 直出为主，R-6）。

两个入口（共享帧校验与 GIF 合成核心）：
- build_animated_gif(request)  -> bytes：POST /api/generate-gif，直接返回 image/gif；
- build_animated_pptx(request) -> bytes：POST /api/generate-animated-pptx，
  把同一 GIF 以 add_picture 等比居中嵌入 16:9 空白页（v1.11，保留）。

入参 {title?, frames: string[], frame_ms?, transparent?}；frames 为前端离屏
ECharts 采帧得到的 base64 PNG dataURL 数组（≥2 帧）。

背景模式（v1.12）：
- transparent=false（默认）：帧转 RGB 实底（前端已按 UI 主题铺深色/浅色底），
  GIF 不支持半透明，透明像素若不铺底会被填黑；
- transparent=true：保留 RGBA 合成，GIF 仅二元透明（P 模式 + transparency 索引、
  disposal=2 防帧残影）；抗锯齿边缘/半透明填充会被硬切，且文字色随主题，仅适合
  贴到与主题同色的背景。

位图产物、不可编辑，属独立动效产物，不影响 C-1（C-1 仅约束标准导出端点）。
契约错误抛 PptxValidationError("VALIDATION_FAILED")；合成/构建异常抛 ParseError("PPTX_BUILD_FAILED")。
"""
from __future__ import annotations

import base64
import binascii
import io
from typing import Any

from pptx import Presentation
from pptx.util import Inches

from app.modules.parser import ParseError
from app.modules.pptx_builder import PptxValidationError, _add_title

try:  # Pillow 为运行期依赖（requirements.txt）
    from PIL import Image
except ImportError:  # pragma: no cover - 环境缺失时给明确错误
    Image = None  # type: ignore[assignment]

_PNG_PREFIX = "data:image/png;base64,"
DEFAULT_FRAME_MS = 100
_MIN_FRAME_MS = 20
_MAX_FRAME_MS = 2000
# 16:9 版式中的可用画幅（与标准导出图表区一致）
_AREA_WIDTH_IN = 12.3
_AREA_HEIGHT_IN = 5.6
_AREA_HEIGHT_NO_TITLE_IN = 6.5


def _decode_frame(raw: Any, index: int):
    """把一帧 base64 PNG dataURL 解码为 PIL Image（RGBA，保留 alpha）。"""
    if Image is None:  # pragma: no cover
        raise ParseError("PPTX_BUILD_FAILED", "缺少 Pillow 依赖，无法合成 GIF")
    if not isinstance(raw, str) or not raw.startswith(_PNG_PREFIX):
        raise PptxValidationError(
            "VALIDATION_FAILED",
            f"frames[{index}] 必须是 data:image/png;base64, 前缀的 dataURL",
        )
    payload = raw[len(_PNG_PREFIX):]
    try:
        data = base64.b64decode(payload, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise PptxValidationError(
            "VALIDATION_FAILED", f"frames[{index}] base64 解码失败"
        ) from exc
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception as exc:  # noqa: BLE001 - PIL 异常类型众多，统一转契约错误
        raise PptxValidationError(
            "VALIDATION_FAILED", f"frames[{index}] 不是合法 PNG 图片"
        ) from exc
    return img.convert("RGBA")


def _parse_request(request: dict[str, Any]) -> dict[str, Any]:
    """入参校验 + 解码（SPEC §5.6a，不信任前端）。返回 RGBA 帧序列。"""
    if not isinstance(request, dict):
        raise PptxValidationError(
            "VALIDATION_FAILED",
            "请求体必须为 JSON 对象 {title?, frames, frame_ms?, transparent?}",
        )
    frames = request.get("frames")
    if not isinstance(frames, list) or len(frames) < 2:
        raise PptxValidationError("VALIDATION_FAILED", "frames 必须为至少 2 帧的数组")
    frame_ms = request.get("frame_ms", DEFAULT_FRAME_MS)
    if not isinstance(frame_ms, int) or isinstance(frame_ms, bool) or not (
        _MIN_FRAME_MS <= frame_ms <= _MAX_FRAME_MS
    ):
        raise PptxValidationError(
            "VALIDATION_FAILED",
            f"frame_ms 必须为 {_MIN_FRAME_MS}~{_MAX_FRAME_MS} 的整数",
        )
    transparent = request.get("transparent", False)
    if not isinstance(transparent, bool):
        raise PptxValidationError("VALIDATION_FAILED", "transparent 必须为布尔值")
    title = request.get("title")
    if title is not None and not isinstance(title, str):
        raise PptxValidationError("VALIDATION_FAILED", "title 必须为字符串或 null")
    images = [_decode_frame(raw, i) for i, raw in enumerate(frames)]
    return {
        "title": title or "",
        "images_rgba": images,
        "frame_ms": frame_ms,
        "transparent": transparent,
    }


def _compose_gif(images_rgba: list, frame_ms: int, transparent: bool) -> bytes:
    """多帧合成循环动态 GIF（loop=0 无限循环）。

    transparent=False：转 RGB 实底（前端已铺主题底色）。
    transparent=True ：保留 alpha，GIF 二元透明（disposal=2 防帧间残影）。
    """
    if Image is None:  # pragma: no cover
        raise ParseError("PPTX_BUILD_FAILED", "缺少 Pillow 依赖，无法合成 GIF")
    try:
        if transparent:
            frames_out = images_rgba  # RGBA；save(GIF) 自动量化并保留透明索引
            kwargs: dict[str, Any] = {"disposal": 2}
        else:
            frames_out = [im.convert("RGB") for im in images_rgba]
            kwargs = {"optimize": True}
        buf = io.BytesIO()
        frames_out[0].save(
            buf,
            format="GIF",
            save_all=True,
            append_images=frames_out[1:],
            duration=frame_ms,
            loop=0,
            **kwargs,
        )
        return buf.getvalue()
    except ParseError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ParseError("PPTX_BUILD_FAILED", f"GIF 合成失败：{exc}") from exc


def build_animated_gif(request: dict[str, Any]) -> bytes:
    """AnimatedFramesRequest -> .gif 字节流（SPEC §5.6a/§6a，v1.12 主形态）。"""
    spec = _parse_request(request)
    return _compose_gif(spec["images_rgba"], spec["frame_ms"], spec["transparent"])


def build_animated_pptx(request: dict[str, Any]) -> bytes:
    """AnimatedFramesRequest -> .pptx 字节流（SPEC §5.6a/§6a，v1.11 保留）。"""
    spec = _parse_request(request)
    # 嵌入 pptx 时透明 GIF 叠在白色版式上文字可能不可见，但遵循用户显式选择。
    gif_bytes = _compose_gif(spec["images_rgba"], spec["frame_ms"], spec["transparent"])

    try:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _add_title(slide, spec["title"])

        # 等比缩放 + 水平垂直居中嵌入（画幅与标准导出图表区一致）
        img_w, img_h = spec["images_rgba"][0].size
        area_w = _AREA_WIDTH_IN
        area_h = _AREA_HEIGHT_IN if spec["title"] else _AREA_HEIGHT_NO_TITLE_IN
        top_edge = 1.2 if spec["title"] else 0.5
        scale = min(area_w / img_w, area_h / img_h)
        w_in, h_in = img_w * scale, img_h * scale
        left = Inches(0.5 + (area_w - w_in) / 2)
        top = Inches(top_edge + (area_h - h_in) / 2)
        slide.shapes.add_picture(
            io.BytesIO(gif_bytes), left, top, Inches(w_in), Inches(h_in)
        )

        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()
    except ParseError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ParseError("PPTX_BUILD_FAILED", f"动效版 PPTX 构建失败：{exc}") from exc
