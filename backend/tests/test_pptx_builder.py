"""M4 pptx_builder 测试（SPEC §6 M4 验收 A-4.x / §5.6）。

覆盖用户要求的三条断言：
1. 所有 has_chart shape 的 chart_data.IsLinked 均为 False（无外部链接）；
2. combo 请求生成的对象同时含 barChart + lineChart + 副轴；
3. replace_data 修改内嵌数据后，图表可更新（读回验证并重新落盘）。

另覆盖：null → dispBlanksAs=gap、副轴 0.0% / 主轴「亿」、非法请求 400、HTTP 端点冒烟。
"""
from __future__ import annotations

import io
import zipfile

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from fastapi.testclient import TestClient

from app.main import app
from app.modules.pptx_builder import PptxValidationError, build_pptx


COMBO_REQUEST = {
    "chart_type": "combo",
    "title": "规模与同比",
    "config": {"dimension": "平台"},
    "data": {
        "categories": ["微信", "抖音", "快手", "微博"],
        "series": [
            {"name": "2025 MAU", "kind": "bar", "axis": "primary", "unit": "yi",
             "data": [14.2, 11.7, 7.0, 5.7]},
            {"name": "同比", "kind": "line", "axis": "secondary", "unit": "percent",
             "data": [0.024, 0.112, -0.028, -0.039]},
        ],
    },
}

LINE_REQUEST = {
    "chart_type": "line",
    "title": "趋势",
    "config": {"dimension": "平台"},
    "data": {
        "categories": ["微信", "抖音", "快手"],
        "series": [
            {"name": "2024 DAU", "kind": "line", "unit": "yi",
             "data": [None, 7.0, 4.1]},
        ],
    },
}


def _chart_xml(blob: bytes) -> str:
    """从 pptx 中取出唯一 chart XML（测试请求均为单图）。"""
    with zipfile.ZipFile(io.BytesIO(blob)) as zf:
        names = [n for n in zf.namelist() if n.startswith("ppt/charts/chart") and n.endswith(".xml")]
        assert names, "pptx 内未找到 chart XML"
        return zf.read(names[0]).decode("utf-8")


def _charts(blob: bytes):
    prs = Presentation(io.BytesIO(blob))
    slide = prs.slides[0]
    return prs, slide, [sh.chart for sh in slide.shapes if sh.has_chart]


# ---------- 断言 1：全部图表 IsLinked=False（无「外部」链接，内嵌 xlsx 保留） ----------

def _external_link_targets(blob: bytes) -> list[str]:
    """收集 chart part 指向外部文件（TargetMode=External）的关系，即 IsLinked=True。

    指向内部 ../embeddings/*.xlsx 的 package 关系是内嵌工作簿，IsLinked=False，
    且是 WPS 双击编辑数据的依据，必须保留。
    """
    prs = Presentation(io.BytesIO(blob))
    external: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_chart:
                continue
            for rel in shape.chart.part.rels.values():
                if rel.is_external:
                    external.append(rel.target_ref)
    return external


def test_all_charts_unlinked_combo():
    blob = build_pptx(COMBO_REQUEST)
    assert _external_link_targets(blob) == []
    # 内嵌工作簿仍在（可双击编辑），且不含指向包外文件的链接
    _, _, charts = _charts(blob)
    assert charts, "至少存在 1 个图表"
    xml = _chart_xml(blob)
    assert "TargetMode=\"External\"" not in xml


def test_all_charts_unlinked_line():
    blob = build_pptx(LINE_REQUEST)
    assert _external_link_targets(blob) == []


# ---------- 断言 2：combo 对象含 barChart + lineChart + 副轴 ----------

def test_combo_chart_structure():
    blob = build_pptx(COMBO_REQUEST)
    xml = _chart_xml(blob)
    assert "<c:barChart>" in xml
    assert "<c:lineChart>" in xml
    # 双值轴：主轴 + 次轴（副轴 axPos=left 是 ablechart 的表达方式，共两个 c:valAx）
    assert xml.count("<c:valAx>") == 2
    # 副轴百分比格式 0.0%
    assert 'formatCode="0.0%"' in xml
    # 主轴「亿」自定义格式
    assert '"0.0\\"' in xml or '亿' in xml
    _, _, charts = _charts(blob)
    assert len(charts) == 1
    # 两个系列：柱 + 线
    assert len(charts[0].series) == 2


# ---------- 断言 3：replace_data 修改内嵌数据后图表可更新 ----------

def test_replace_data_updates_chart_and_persists():
    blob = build_pptx(COMBO_REQUEST)
    prs, slide, _ = _charts(blob)
    chart = next(sh.chart for sh in slide.shapes if sh.has_chart)
    assert list(chart.series[0].values)[:4] == [14.2, 11.7, 7.0, 5.7]

    new_data = CategoryChartData()
    new_data.categories = ["微信", "抖音", "快手", "微博"]
    new_data.add_series("2025 MAU", (20.0, 18.0, 9.0, 6.0))
    new_data.add_series("同比", (0.05, 0.20, -0.01, -0.08))
    chart.replace_data(new_data)

    # 内存对象立即反映
    assert list(chart.series[0].values) == [20.0, 18.0, 9.0, 6.0]

    # 重新落盘再读回，确认内嵌 workbook 与图表同步更新
    out = io.BytesIO()
    prs.save(out)
    _, _, reopened = _charts(out.getvalue())
    assert list(reopened[0].series[0].values) == [20.0, 18.0, 9.0, 6.0]
    assert list(reopened[0].series[1].values) == [0.05, 0.20, -0.01, -0.08]


# ---------- C-3：null 断线 gap，不把空值当 0 ----------

def test_line_null_becomes_gap_not_zero():
    blob = build_pptx(LINE_REQUEST)
    xml = _chart_xml(blob)
    assert '<c:dispBlanksAs val="gap"/>' in xml
    _, _, charts = _charts(blob)
    values = list(charts[0].series[0].values)
    assert values[0] is None
    assert values[1] == 7.0
    assert 0 not in [v for v in values if v is not None]


# ---------- 纯 bar / 折线走 add_chart ----------

def test_bar_uses_native_chart():
    req = {
        "chart_type": "bar",
        "config": {"dimension": "平台"},
        "data": {
            "categories": ["a", "b"],
            "series": [{"name": "v", "kind": "bar", "unit": "yi", "data": [1.0, 2.0]}],
        },
    }
    blob = build_pptx(req)
    xml = _chart_xml(blob)
    assert "<c:barChart>" in xml
    assert "<c:lineChart>" not in xml
    assert _external_link_targets(blob) == []


# ---------- 校验 ----------

def test_rejects_unknown_chart_type():
    with pytest.raises(PptxValidationError):
        build_pptx({"chart_type": "radar", "data": COMBO_REQUEST["data"]})


def test_rejects_combo_without_secondary():
    bad = {
        "chart_type": "combo",
        "data": {"categories": ["a"], "series": [{"name": "v", "kind": "bar", "data": [1.0]}]},
    }
    with pytest.raises(PptxValidationError):
        build_pptx(bad)


# ---------- HTTP 端点冒烟：返回 .pptx 文件流 ----------

def test_generate_pptx_endpoint_returns_file():
    client = TestClient(app)
    resp = client.post("/api/generate-pptx", json=COMBO_REQUEST)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert ".pptx" in resp.headers.get("content-disposition", "")
    # 魔数 PK\x03\x04
    assert resp.content[:2] == b"PK"
    # 产物可被 python-pptx 正常打开
    prs = Presentation(io.BytesIO(resp.content))
    assert sum(1 for sh in prs.slides[0].shapes if sh.has_chart) == 1


def test_generate_pptx_endpoint_bad_request():
    client = TestClient(app)
    # scatter 属六种受识别图型，但 M4 尚未实现 → UNSUPPORTED_CHART/400
    resp = client.post("/api/generate-pptx", json={"chart_type": "scatter", "data": {}})
    assert resp.status_code == 400
    assert resp.json()["code"] == "UNSUPPORTED_CHART"
    # 完全未知的图型 → VALIDATION_FAILED/400
    resp2 = client.post("/api/generate-pptx", json={"chart_type": "radar", "data": {}})
    assert resp2.status_code == 400
    assert resp2.json()["code"] == "VALIDATION_FAILED"
