"""M4b 动效版 GIF 导出测试（SPEC §5.6a/§6a，v1.11，R-6）。

覆盖：
1. 合法帧序列 → PK 魔数 .pptx、内含 GIF 媒体（GIF89a）、slide 含图片 shape；
2. GIF 为循环动图（n_frames 与输入一致、loop=0）；
3. 非法入参 → VALIDATION_FAILED（400）；
4. 端点冒烟 → chart-animated 文件名 + 正确 MIME。
"""
from __future__ import annotations

import base64
import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from PIL import Image

from app.main import app
from app.modules.gif_builder import build_animated_pptx
from app.modules.pptx_builder import PptxValidationError


def _frame(color: tuple[int, int, int], size: tuple[int, int] = (320, 200)) -> str:
    """生成一帧纯色 PNG dataURL。"""
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode("ascii")


FRAMES = [_frame((200, 30, 30)), _frame((30, 200, 30)), _frame((30, 30, 200))]


def _gif_part(blob: bytes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(blob)) as zf:
        names = [n for n in zf.namelist() if n.startswith("ppt/media/") and n.endswith(".gif")]
        assert names, "pptx 中应包含 GIF 媒体"
        return zf.read(names[0])


def test_build_animated_pptx_contains_looping_gif():
    blob = build_animated_pptx({"title": "动效演示", "frames": FRAMES, "frame_ms": 120})
    assert blob[:2] == b"PK"
    gif = _gif_part(blob)
    assert gif[:6] == b"GIF89a"
    img = Image.open(io.BytesIO(gif))
    assert getattr(img, "n_frames", 1) == len(FRAMES)
    assert img.info.get("loop") == 0  # 无限循环
    assert img.info.get("duration") == 120


def test_slide_has_picture_shape():
    from pptx import Presentation

    blob = build_animated_pptx({"frames": FRAMES})
    prs = Presentation(io.BytesIO(blob))
    slide = prs.slides[0]
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]  # PICTURE
    assert len(pics) == 1


@pytest.mark.parametrize(
    "payload",
    [
        {},                                            # 缺 frames
        {"frames": []},                                # 空数组
        {"frames": [FRAMES[0]]},                       # 少于 2 帧
        {"frames": ["not-a-dataurl", FRAMES[0]]},      # 非 dataURL
        {"frames": ["data:image/png;base64,%%%", FRAMES[0]]},  # base64 非法
        {"frames": FRAMES, "frame_ms": 5},             # frame_ms 越界
        {"frames": FRAMES, "title": 123},              # title 非字符串
    ],
)
def test_invalid_request_raises(payload):
    with pytest.raises(PptxValidationError):
        build_animated_pptx(payload)


def test_endpoint_smoke():
    client = TestClient(app)
    resp = client.post(
        "/api/generate-animated-pptx",
        json={"title": "动效演示", "frames": FRAMES, "frame_ms": 100},
    )
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument")
    disposition = resp.headers.get("content-disposition", "")
    assert "chart-animated" in disposition and disposition.endswith('.pptx"')


def test_endpoint_invalid_returns_400():
    client = TestClient(app)
    resp = client.post("/api/generate-animated-pptx", json={"frames": []})
    assert resp.status_code == 400
    assert resp.json()["code"] == "VALIDATION_FAILED"
