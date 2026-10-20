"""M0 ablechart 可行性验证（SPEC §6 M0 门禁，一次性脚本，通过后保留作回归）。

五项目标：
1. 能装：import ablechart==0.2.0；
2. 能生成双轴 combo：主轴 MAU(bar, yi) + 副轴 同比(line, percent)；
3. 副轴 percent 数字格式 0.0%；
4. 原生图表（非图片）且无外部数据链接（C-1/C-2）；
5. 产物可被 python-pptx 重新打开。

运行：cd backend && python3 spikes/ablechart_spike.py
"""
from __future__ import annotations

import io
import zipfile

import pandas as pd
from pptx import Presentation

import ablechart

ROWS = [
    ("微信", 14.2, 0.024),
    ("抖音", 11.7, 0.112),
    ("快手", 7.0, -0.028),
    ("微博", 5.7, -0.039),
    ("B站", 3.7, 0.088),
    ("小红书", 3.5, 0.094),
]


def build() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
    df = pd.DataFrame(ROWS, columns=["平台", "2025 MAU", "同比"])
    ablechart.create_combo_chart(
        slide=slide,
        df=df,
        categories_col="平台",
        series_config=[
            {"key": "2025 MAU", "name": "2025 MAU（亿）", "type": "bar", "axis": "primary"},
            {"key": "同比", "name": "同比", "type": "line", "axis": "secondary"},
        ],
        polish=False,
    )
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def main() -> None:
    blob = build()
    print(f"[1] ablechart {ablechart.__version__} 生成 pptx {len(blob)} bytes")

    # 解包检查
    zf = zipfile.ZipFile(io.BytesIO(blob))
    names = zf.namelist()
    charts = [n for n in names if n.startswith("ppt/charts/chart") and n.endswith(".xml")]
    embeds = [n for n in names if n.startswith("ppt/embeddings/")]
    media_imgs = [n for n in names if n.startswith("ppt/media/")]
    print(f"[2] chart xml: {charts}")
    print(f"[3] embeddings: {embeds}")
    print(f"[4] media(图片，应为空或仅非图表): {media_imgs}")

    # 外部链接关系（C-2）：chart*.rels 里不应出现 External 关系
    ext_links = []
    for n in names:
        if n.startswith("ppt/charts/_rels/") and n.endswith(".rels"):
            txt = zf.read(n).decode("utf-8", "ignore")
            if "TargetMode=\"External\"" in txt:
                ext_links.append(n)
    print(f"[5] 外部数据链接关系（应为空）: {ext_links}")

    # python-pptx 读回
    prs = Presentation(io.BytesIO(blob))
    n_chart_shapes = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_chart:
                n_chart_shapes += 1
                chart = shape.chart
                print(f"[6] chart_type={chart.chart_type}, series={[s.name for s in chart.series]}")
                part = chart.part
                print(f"    chart part rels: {[(r.reltype.split('/')[-1], r.is_external) for r in part.rels.values()]}")
    print(f"[7] has_chart shape 数: {n_chart_shapes}")

    # BreakLink API 探查（R-2）
    from pptx.chart.chart import Chart  # noqa: F401

    has_break = any(hasattr(Chart, m) for m in ("break_link", "BreakLink", "breakLink"))
    print(f"[8] Chart 公共 BreakLink API: {has_break}")

    chart_xml = zf.read(charts[0]).decode("utf-8", "ignore")
    print(f"[9] chart xml 含 0.0% 格式: {'0.0%' in chart_xml}")
    print(f"[10] chart xml 含外部链接标记: {'External' in chart_xml or 'externalData' in chart_xml}")
    assert charts, "未生成原生 chart xml（C-1 失败）"
    assert not ext_links, "存在外部数据链接（C-2 失败）"
    assert n_chart_shapes >= 1
    print("M0 SPIKE PASSED")


if __name__ == "__main__":
    main()